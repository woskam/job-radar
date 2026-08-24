"""
Renders the structured short-CV data (letters/cv_short.yaml / cv_short_nl.yaml)
into a one-page docx/pdf resume with its own distinct layout -- a top-right
tagline/contact block, an italic pitch line, right-aligned dates per entry,
inline **bold**/[link](url) spans inside bullets and descriptions, a
left-accent-bar projects block, and a label/value skills table. Deliberately
a separate module and layout from letters/cv_builder.py (the long CV): this
design isn't a shorter version of that template, it has different layout
primitives entirely (right-aligned dates, inline markup, a table) that the
long CV's plain-text block parser can't represent. Reuses the low-level
shared bits from document_style.py (SENDER, MUTED_RGB, docx_add_border, the
inline-markup renderers) -- not its letter-specific header/footer builders.
"""

import sys
from pathlib import Path

import yaml

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))

from letters.document_style import (
    MUTED_RGB,
    SENDER,
    docx_add_border,
    docx_add_inline_runs,
    parse_inline_markup,
    pdf_write_inline,
)

CV_SHORT_PATH_EN = ROOT / "letters" / "cv_short.yaml"
CV_SHORT_PATH_NL = ROOT / "letters" / "cv_short_nl.yaml"

SECTION_TITLES = {
    "en": {"experience": "EXPERIENCE", "projects": "SELECTED PROJECTS", "skills": "SKILLS & EDUCATION"},
    "nl": {"experience": "ERVARING", "projects": "GESELECTEERDE PROJECTEN", "skills": "VAARDIGHEDEN & OPLEIDING"},
}

# Approximates the dashboard's own accent blue (oklch(0.55 0.09 215)) for the
# projects accent bar -- close enough for a small decorative rule, not meant
# to be a pixel-exact color match.
ACCENT_RGB = (90, 120, 150)


def _city_from_address(address: str) -> str:
    # "Some Street 5, 1234 AB Amsterdam" -> "Amsterdam" -- the short CV
    # shows just the city, not the full street address.
    last_segment = address.split(",")[-1].strip()
    return last_segment.split()[-1] if last_segment else address


def load_short_cv_data(language: str = "en") -> dict:
    path = CV_SHORT_PATH_NL if language == "nl" else CV_SHORT_PATH_EN
    if not path.exists() and path == CV_SHORT_PATH_NL:
        path = CV_SHORT_PATH_EN
    with open(path) as f:
        return yaml.safe_load(f)


def _pdf_safe(text: str) -> str:
    # Core PDF fonts only support cp1252 -- "→" has no equivalent there
    # (unlike em/en dashes and the bullet, which cp1252 does support), so
    # sub it and fall back to a lossy replace for anything else unexpected.
    text = text.replace("→", "->")
    return text.encode("cp1252", "replace").decode("cp1252")


def _sanitize_for_pdf(value):
    if isinstance(value, str):
        return _pdf_safe(value)
    if isinstance(value, list):
        return [_sanitize_for_pdf(v) for v in value]
    if isinstance(value, dict):
        return {k: _sanitize_for_pdf(v) for k, v in value.items()}
    return value


def _strip_table_borders(table) -> None:
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    tbl_pr = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "nil")
        borders.append(el)
    tbl_pr.append(borders)


def build_short_cv_docx(language: str = "en", font_name: str = "Arial"):
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
    from docx.shared import Pt, RGBColor

    data = load_short_cv_data(language)
    titles = SECTION_TITLES.get(language, SECTION_TITLES["en"])
    muted = RGBColor(*MUTED_RGB)
    accent = RGBColor(*ACCENT_RGB)

    doc = Document()
    doc.styles["Normal"].font.name = font_name
    doc.styles["Normal"].font.size = Pt(10)
    section = doc.sections[0]
    usable_width = section.page_width - section.left_margin - section.right_margin

    # Header: name on the left, tagline + contact stacked and right-aligned,
    # both on the same visual row -- a borderless 2-column table, the
    # standard python-docx idiom for "two blocks on one line".
    header = doc.add_table(rows=1, cols=2)
    header.autofit = False
    header.columns[0].width = int(usable_width * 0.55)
    header.columns[1].width = usable_width - int(usable_width * 0.55)
    _strip_table_borders(header)

    name_run = header.cell(0, 0).paragraphs[0].add_run(SENDER["name"].upper())
    name_run.bold = True
    name_run.font.size = Pt(24)
    name_run.font.name = font_name

    contact_cell = header.cell(0, 1)
    contact_lines = [data["tagline"], f"{_city_from_address(SENDER['address'])} · {SENDER['email']}", SENDER["url"]]
    for i, line in enumerate(contact_lines):
        p = contact_cell.paragraphs[0] if i == 0 else contact_cell.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        run = p.add_run(line)
        run.font.name = font_name
        run.font.size = Pt(8.5)
        run.font.color.rgb = muted
        if i == 0:
            run.bold = True

    # Pitch line: plain text, last clause in italic.
    pitch_p = doc.add_paragraph()
    pitch_p.paragraph_format.space_before = Pt(10)
    pitch_run = pitch_p.add_run(data["pitch_plain"])
    pitch_run.font.name = font_name
    pitch_run.font.size = Pt(10.5)
    italic_run = pitch_p.add_run(data["pitch_italic"])
    italic_run.italic = True
    italic_run.font.name = font_name
    italic_run.font.size = Pt(10.5)
    docx_add_border(pitch_p, "bottom")

    def add_section_heading(text: str) -> None:
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(12)
        p.paragraph_format.space_after = Pt(4)
        run = p.add_run(text)
        run.bold = True
        run.font.name = font_name
        run.font.size = Pt(11)
        run.font.color.rgb = muted

    def add_entry_line(title: str, company: str, dates: str) -> None:
        p = doc.add_paragraph()
        p.paragraph_format.tab_stops.add_tab_stop(usable_width, WD_TAB_ALIGNMENT.RIGHT)
        title_run = p.add_run(f"{title} · {company}")
        title_run.bold = True
        title_run.font.name = font_name
        title_run.font.size = Pt(11)
        date_run = p.add_run(f"\t{dates}")
        date_run.font.name = font_name
        date_run.font.size = Pt(9)
        date_run.font.color.rgb = muted

    def add_bullet(text: str) -> None:
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Pt(16)
        p.paragraph_format.first_line_indent = Pt(-16)
        p.paragraph_format.tab_stops.add_tab_stop(Pt(16))
        p.add_run("•\t").font.name = font_name
        docx_add_inline_runs(p, text, font_name, size=10)

    # EXPERIENCE
    add_section_heading(titles["experience"])
    for entry in data["experience"]:
        add_entry_line(entry["title"], entry["company"], entry["dates"])
        for bullet in entry["bullets"]:
            add_bullet(bullet)

    for entry in data.get("compact_experience", []):
        add_entry_line(entry["title"], entry["company"], entry["dates"])
        desc_p = doc.add_paragraph()
        docx_add_inline_runs(desc_p, entry["description"], font_name, size=10)

    # SELECTED PROJECTS
    suffix = data.get("projects_heading_suffix")
    add_section_heading(f"{titles['projects']} — {suffix}" if suffix else titles["projects"])
    for project in data["projects"]:
        name_p = doc.add_paragraph()
        name_p.paragraph_format.left_indent = Pt(12)
        docx_add_border(name_p, "left")
        name_run = name_p.add_run(f"{project['name']}  ")
        name_run.bold = True
        name_run.font.name = font_name
        name_run.font.size = Pt(10.5)
        tech_run = name_p.add_run(project["tech"])
        tech_run.font.name = font_name
        tech_run.font.size = Pt(7.5)
        tech_run.font.color.rgb = muted

        desc_p = doc.add_paragraph()
        desc_p.paragraph_format.left_indent = Pt(12)
        docx_add_border(desc_p, "left")
        docx_add_inline_runs(desc_p, project["description"], font_name, size=9.5)

    # SKILLS & EDUCATION
    add_section_heading(titles["skills"])
    table = doc.add_table(rows=len(data["skills_table"]), cols=2)
    table.autofit = False
    table.columns[0].width = int(usable_width * 0.18)
    table.columns[1].width = usable_width - int(usable_width * 0.18)
    _strip_table_borders(table)
    for row_idx, row in enumerate(data["skills_table"]):
        label_run = table.cell(row_idx, 0).paragraphs[0].add_run(row["label"])
        label_run.bold = True
        label_run.font.name = font_name
        label_run.font.size = Pt(8.5)
        label_run.font.color.rgb = muted

        value_run = table.cell(row_idx, 1).paragraphs[0].add_run(row["text"])
        value_run.font.name = font_name
        value_run.font.size = Pt(9.5)

    return doc


BULLET_INDENT = 5
PROJECT_INDENT = 4


def _wrapped_line_count(pdf, text: str, avail_width: float) -> int:
    import math

    # Measure only the visible text -- "**"/"[...]"/"(url)" markup characters
    # never render, so including them would overestimate the width and the
    # resulting line/page-break math.
    visible = "".join(frag["text"] for frag in parse_inline_markup(text))
    width = pdf.get_string_width(visible)
    return max(1, math.ceil(width / avail_width))


def _ensure_space(pdf, needed_height: float) -> None:
    # Pre-empt fpdf's own automatic page break: an auto-break that fires in
    # the middle of a label+value (or name+description) pair strands the
    # second half at a stale y-coordinate captured before the break, since
    # only the first half's page actually changed. Checking first and
    # breaking the whole unit ourselves avoids that class of bug entirely.
    if pdf.get_y() + needed_height > pdf.h - pdf.b_margin:
        pdf.add_page()


def build_short_cv_pdf(language: str = "en", font_family: str = "Helvetica"):
    from fpdf import FPDF
    from fpdf.enums import XPos, YPos

    data = _sanitize_for_pdf(load_short_cv_data(language))
    titles = SECTION_TITLES.get(language, SECTION_TITLES["en"])

    pdf = FPDF()
    pdf.core_fonts_encoding = "cp1252"
    pdf.set_margins(left=16, top=10, right=16)
    pdf.set_auto_page_break(True, margin=12)
    pdf.add_page()

    base_margin = pdf.l_margin
    right_edge = pdf.w - pdf.r_margin
    full_width = right_edge - base_margin

    # Header: name on the left, tagline + contact stacked and right-aligned.
    y0 = pdf.get_y()
    pdf.set_font(font_family, "B", 20)
    pdf.set_text_color(0, 0, 0)
    pdf.set_xy(base_margin, y0)
    pdf.cell(full_width, 9, SENDER["name"].upper())

    pdf.set_font(font_family, "B", 8)
    pdf.set_text_color(*MUTED_RGB)
    pdf.set_xy(base_margin, y0)
    pdf.cell(full_width, 4.5, data["tagline"], align="R")

    pdf.set_font(font_family, "", 8)
    pdf.set_xy(base_margin, y0 + 4.5)
    city = _city_from_address(SENDER["address"])
    pdf.cell(full_width, 4.5, f"{city} · {SENDER['email']}", align="R")
    pdf.set_xy(base_margin, y0 + 9)
    pdf.cell(full_width, 4.5, SENDER["url"], align="R")

    pdf.set_y(y0 + 13)
    pdf.set_x(base_margin)
    pdf.set_font(font_family, "", 10)
    pdf.set_text_color(0, 0, 0)
    pdf.write(5.2, data["pitch_plain"])
    pdf.set_font(font_family, "I", 10)
    pdf.write(5.2, data["pitch_italic"])
    pdf.ln(6)

    pdf.set_draw_color(*MUTED_RGB)
    pdf.set_line_width(0.3)
    rule_y = pdf.get_y()
    pdf.line(base_margin, rule_y, right_edge, rule_y)
    pdf.ln(3)

    HEADING_ROW_HEIGHT = 5.2
    HEADING_GAP = 1.2

    def section_heading(text: str, next_height: float = 0) -> None:
        # Keep the heading with at least its first entry -- a heading alone
        # at the bottom of a page with its content pushed to the next one
        # looks broken (same fix as the long CV's cv_builder.py).
        _ensure_space(pdf, HEADING_ROW_HEIGHT + HEADING_GAP + next_height)
        pdf.set_font(font_family, "B", 10.5)
        pdf.set_text_color(*MUTED_RGB)
        pdf.cell(0, HEADING_ROW_HEIGHT, text, new_x=XPos.LMARGIN, new_y=YPos.NEXT)
        pdf.ln(HEADING_GAP)

    def entry_line(title: str, company: str, dates: str) -> None:
        pdf.set_text_color(0, 0, 0)
        pdf.set_font(font_family, "B", 10)
        title_text = f"{title} · {company}"
        title_width = pdf.get_string_width(title_text)
        pdf.cell(title_width, 5.2, title_text)
        pdf.set_font(font_family, "", 8)
        pdf.set_text_color(*MUTED_RGB)
        pdf.cell(0, 5.2, dates, align="R", new_x=XPos.LMARGIN, new_y=YPos.NEXT)
        pdf.set_text_color(0, 0, 0)

    def bullet_line(text: str) -> None:
        pdf.set_font(font_family, "", 9.5)
        pdf.set_text_color(0, 0, 0)
        pdf.cell(BULLET_INDENT, 4.8, "•")
        pdf.set_left_margin(base_margin + BULLET_INDENT)
        pdf_write_inline(pdf, text, font_family, 9.5, 4.8)
        pdf.ln(4.8)
        pdf.set_left_margin(base_margin)

    def entry_height(entry: dict) -> float:
        # get_string_width() measures using whatever font is currently active
        # on the pdf object, so it must be set to match the font this text
        # actually renders in before measuring -- otherwise the estimate
        # silently uses whatever font a previous section left active.
        pdf.set_font(font_family, "", 9.5)
        bullets_height = sum(
            _wrapped_line_count(pdf, b, full_width - BULLET_INDENT) * 4.8 for b in entry["bullets"]
        )
        return 5.2 + bullets_height + 1.2

    def compact_entry_height(entry: dict) -> float:
        pdf.set_font(font_family, "", 9)
        return 5.2 + _wrapped_line_count(pdf, entry["description"], full_width) * 4.6 + 1.2

    def project_height(project: dict) -> float:
        pdf.set_font(font_family, "", 7.5)
        tech_lines = _wrapped_line_count(pdf, project["tech"], full_width - PROJECT_INDENT)
        pdf.set_font(font_family, "", 9)
        desc_lines = _wrapped_line_count(pdf, project["description"], full_width - PROJECT_INDENT)
        return tech_lines * 4.8 + desc_lines * 4.6 + 1.5

    def row_height(row: dict) -> float:
        pdf.set_font(font_family, "", 9)
        return _wrapped_line_count(pdf, row["text"], full_width - label_width) * 4.6 + 1.0

    # A short CV's content rarely lines up exactly with a full page, and a
    # fixed, compact spacing (tuned to *guarantee* everything fits even for a
    # longer version of this CV) leaves a big dead zone at the bottom for a
    # shorter one -- exactly what was flagged as looking sparse compared to
    # the reference PDF. Rather than hand-tune the constants above for
    # whatever the CV happens to contain today (which just breaks again the
    # next time a bullet's added or removed), measure the natural height of
    # everything below the header once, compare it to the space actually
    # available on the page, and spread the leftover evenly across the gaps
    # between entries/projects/rows -- same idea as "justify", vertically.
    label_width = 28
    content_top_y = pdf.get_y()
    natural_height = (
        3 * (HEADING_ROW_HEIGHT + HEADING_GAP)
        + sum(entry_height(e) for e in data["experience"])
        + sum(compact_entry_height(e) for e in data.get("compact_experience", []))
        + sum(project_height(p) for p in data["projects"])
        + sum(row_height(r) for r in data["skills_table"])
    )
    gap_points = (
        len(data["experience"]) + len(data.get("compact_experience", [])) + len(data["projects"])
        + len(data["skills_table"])
    )
    available_height = (pdf.h - pdf.b_margin) - content_top_y
    slack = available_height - natural_height
    # 92% safety margin against the estimate being slightly optimistic (font
    # metrics/wrapping approximations), capped so a very short CV doesn't end
    # up with absurdly loose spacing instead of just... having a bit of air.
    extra_gap = max(0.0, min(slack / gap_points * 0.92, 4.0)) if gap_points else 0.0

    # EXPERIENCE
    first_entry_height = entry_height(data["experience"][0]) if data["experience"] else 0
    section_heading(titles["experience"], next_height=first_entry_height)
    for entry in data["experience"]:
        _ensure_space(pdf, entry_height(entry))
        entry_line(entry["title"], entry["company"], entry["dates"])
        for bullet in entry["bullets"]:
            bullet_line(bullet)
        pdf.ln(1.2 + extra_gap)

    for entry in data.get("compact_experience", []):
        pdf.set_font(font_family, "", 9)
        desc_height = _wrapped_line_count(pdf, entry["description"], full_width) * 4.6
        _ensure_space(pdf, 5.2 + desc_height + 1.2)
        entry_line(entry["title"], entry["company"], entry["dates"])
        pdf.set_text_color(0, 0, 0)
        pdf_write_inline(pdf, entry["description"], font_family, 9, 4.6)
        pdf.ln(4.6)
        pdf.ln(1.2 + extra_gap)

    # SELECTED PROJECTS
    suffix = data.get("projects_heading_suffix")
    first_project_height = project_height(data["projects"][0]) if data["projects"] else 0
    section_heading(
        f"{titles['projects']} — {suffix}" if suffix else titles["projects"], next_height=first_project_height
    )
    for project in data["projects"]:
        _ensure_space(pdf, project_height(project))

        y_start = pdf.get_y()
        pdf.set_left_margin(base_margin + PROJECT_INDENT)
        pdf.set_x(base_margin + PROJECT_INDENT)

        pdf.set_font(font_family, "B", 10)
        pdf.set_text_color(0, 0, 0)
        name_text = f"{project['name']}  "
        name_width = pdf.get_string_width(name_text)
        pdf.cell(name_width, 4.8, name_text)
        pdf.set_font(font_family, "", 7.5)
        pdf.set_text_color(*MUTED_RGB)
        pdf.multi_cell(0, 4.8, project["tech"], align="L", new_x=XPos.LMARGIN, new_y=YPos.NEXT)

        pdf.set_text_color(0, 0, 0)
        pdf_write_inline(pdf, project["description"], font_family, 9, 4.6)
        pdf.ln(4.6)

        pdf.set_left_margin(base_margin)
        y_end = pdf.get_y()
        pdf.set_draw_color(*ACCENT_RGB)
        pdf.set_line_width(0.7)
        pdf.line(base_margin + 0.5, y_start + 1, base_margin + 0.5, y_end - 1.5)
        pdf.ln(1.5 + extra_gap)

    # SKILLS & EDUCATION
    first_row_height = row_height(data["skills_table"][0]) if data["skills_table"] else 0
    section_heading(titles["skills"], next_height=first_row_height)
    for row in data["skills_table"]:
        _ensure_space(pdf, row_height(row))

        row_y = pdf.get_y()
        pdf.set_font(font_family, "B", 8)
        pdf.set_text_color(*MUTED_RGB)
        pdf.set_xy(base_margin, row_y)
        pdf.multi_cell(label_width, 4.6, row["label"], align="L")

        pdf.set_left_margin(base_margin + label_width)
        pdf.set_xy(base_margin + label_width, row_y)
        pdf.set_font(font_family, "", 9)
        pdf.set_text_color(0, 0, 0)
        pdf.multi_cell(0, 4.6, row["text"], align="L", new_x=XPos.LMARGIN, new_y=YPos.NEXT)
        pdf.set_left_margin(base_margin)
        pdf.ln(1.0 + extra_gap)

    return pdf
