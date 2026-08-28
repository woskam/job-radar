import io
import os
import re
import sqlite3
import sys
from pathlib import Path

from flask import Flask, redirect, render_template, request, send_file, url_for

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))

from db.migrations import ensure_columns
from letters.cv_builder import build_cv_docx, build_cv_pdf
from letters.cv_short_builder import CV_SHORT_PATH_NL, build_short_cv_docx, build_short_cv_pdf
from letters.document_style import DEFAULT_FONT, FONT_OPTIONS, SENDER
from letters.generator import (
    CV_PATH_NL,
    detect_language,
    generate_letter,
    load_cv,
    load_projects,
    save_letter,
)
from letters.interview_prep import generate_interview_prep, save_interview_prep
from matching.scorer import load_config, score_breakdown

DRY_RUN = os.environ.get("DRY_RUN", "true").lower() == "true"

DB_PATH = Path(os.environ.get("JOB_RADAR_DB_PATH", ROOT / "db" / "job_radar.db"))

app = Flask(__name__)

# Status pill colors, from the design handoff -- extended with 'approved' and
# 'sent', which weren't in the original handoff but do exist here, following
# the same oklch(0.95 .. )/oklch(0.45 ..) formula.
STATUS_META = {
    "pending_approval": {"label": "Needs review", "bg": "#F4F1EA", "fg": "#6E6A63"},
    "approved": {"label": "Writing letter", "bg": "#F4F1EA", "fg": "#6E6A63"},
    "letter_drafted": {"label": "Letter ready", "bg": "oklch(0.95 0.035 95)", "fg": "oklch(0.45 0.08 80)"},
    "reviewed": {"label": "Reviewed", "bg": "oklch(0.95 0.03 215)", "fg": "oklch(0.45 0.09 215)"},
    "sent": {"label": "Sent", "bg": "oklch(0.95 0.05 145)", "fg": "oklch(0.45 0.1 145)"},
    "interview": {"label": "Interview", "bg": "oklch(0.95 0.04 300)", "fg": "oklch(0.45 0.1 300)"},
    "rejected": {"label": "Rejected", "bg": "#FAF3F2", "fg": "#8A5A55"},
    # Not a real value of jobs.status (that stays 'rejected') -- a pseudo-status
    # for filtering/browsing the subset of rejections where rejected_stage is
    # 'applied' or 'interview' (the employer said no, not a self-reject).
    "rejected_by_employer": {"label": "Rejected by employer", "bg": "#FAF3F2", "fg": "#8A5A55"},
}
DEFAULT_STATUS_META = {"label": "Unknown", "bg": "#F4F1EA", "fg": "#6E6A63"}
REJECTED_BY_EMPLOYER_STAGES = ("applied", "interview")


def _font_choice() -> dict:
    return FONT_OPTIONS.get(request.args.get("font", DEFAULT_FONT), FONT_OPTIONS[DEFAULT_FONT])


def _letter_filename(job: dict, extension: str) -> str:
    raw = f"cover_letter_{job['company']}_{job['title']}"
    safe = re.sub(r"[^A-Za-z0-9._-]+", "_", raw).strip("_")[:120]
    return f"{safe}.{extension}"


def status_meta(status: str) -> dict:
    return STATUS_META.get(status, DEFAULT_STATUS_META)


REMOTE_PREFIX_RE = re.compile(r"(?i)^remote[,\s]+(.+)")


def _normalize_city(location: str) -> str:
    # Different platforms format the same city wildly differently -- EY's
    # SuccessFactors appends "NL, 1083 HP", Workday sometimes adds a state/
    # province, Radancy tacks on "+7 more..." for ambiguous multi-location
    # postings -- but the city name itself is consistently the first
    # comma-separated segment across every platform seen in this project.
    # Used only for the location filter's suggestion list: the underlying
    # filter is still a substring match against the raw location, so
    # picking the clean "Amsterdam" suggestion correctly matches every
    # differently-formatted Amsterdam variant, not just one exact string.
    location = location.strip()

    # Remote-first postings (GitLab confirmed live: "Remote, Netherlands",
    # "Remote Ireland") flip this -- the informative part is the
    # country/region *after* "Remote", not the word "Remote" itself. Taking
    # the first comma segment there would collapse every remote posting from
    # every country into one meaningless "Remote" bucket.
    remote_match = REMOTE_PREFIX_RE.match(location)
    if remote_match:
        location = remote_match.group(1)

    # Postings with multiple locations use ";" as a secondary separator
    # (e.g. "Remote, Canada; Remote, United Kingdom") -- keep just the first.
    location = location.split(";")[0].strip()
    return location.split(",")[0].strip()


def get_db() -> sqlite3.Connection:
    conn = sqlite3.connect(DB_PATH)
    conn.row_factory = sqlite3.Row
    ensure_columns(conn)
    return conn


def _sidebar_counts(conn: sqlite3.Connection) -> dict:
    return {
        "total": conn.execute("SELECT COUNT(*) FROM jobs WHERE status != 'new'").fetchone()[0],
        "pending_approval": conn.execute("SELECT COUNT(*) FROM jobs WHERE status = 'pending_approval'").fetchone()[0],
        "letter_drafted": conn.execute("SELECT COUNT(*) FROM jobs WHERE status = 'letter_drafted'").fetchone()[0],
        "sent": conn.execute("SELECT COUNT(*) FROM jobs WHERE status = 'sent'").fetchone()[0],
        "interview": conn.execute("SELECT COUNT(*) FROM jobs WHERE status = 'interview'").fetchone()[0],
        "rejected": conn.execute("SELECT COUNT(*) FROM jobs WHERE status = 'rejected'").fetchone()[0],
        "rejected_by_employer": conn.execute(
            f"SELECT COUNT(*) FROM jobs WHERE status = 'rejected' AND rejected_stage IN "
            f"({','.join('?' for _ in REJECTED_BY_EMPLOYER_STAGES)})",
            REJECTED_BY_EMPLOYER_STAGES,
        ).fetchone()[0],
    }


PAGE_SIZE = 50
SCORE_OPTIONS = ["0.5", "0.6", "0.65", "0.7", "0.75", "0.8", "0.9", "1.0"]
STRONG_MATCH_THRESHOLD = 0.80


@app.route("/")
def index():
    title = request.args.get("title", "").strip()
    company = request.args.get("company", "").strip()
    location = request.args.get("location", "").strip()
    status = request.args.get("status", "").strip()
    min_score = request.args.get("min_score", "").strip()
    try:
        page = max(1, int(request.args.get("page", 1)))
    except ValueError:
        page = 1

    where = ["status != 'new'"]
    params: list = []

    if title:
        where.append("title LIKE ?")
        params.append(f"%{title}%")
    if company:
        where.append("company = ?")
        params.append(company)
    if location:
        where.append("location LIKE ?")
        params.append(f"%{location}%")
    if status == "rejected_by_employer":
        where.append(
            f"status = 'rejected' AND rejected_stage IN ({','.join('?' for _ in REJECTED_BY_EMPLOYER_STAGES)})"
        )
        params.extend(REJECTED_BY_EMPLOYER_STAGES)
    elif status:
        where.append("status = ?")
        params.append(status)
    if min_score:
        try:
            where.append("relevance_score >= ?")
            params.append(float(min_score))
        except ValueError:
            pass

    where_sql = " AND ".join(where)

    conn = get_db()
    total = conn.execute(f"SELECT COUNT(*) FROM jobs WHERE {where_sql}", params).fetchone()[0]
    total_pages = max(1, -(-total // PAGE_SIZE))  # ceil division
    page = min(page, total_pages)
    offset = (page - 1) * PAGE_SIZE

    jobs = conn.execute(
        f"SELECT * FROM jobs WHERE {where_sql} ORDER BY relevance_score DESC LIMIT ? OFFSET ?",
        params + [PAGE_SIZE, offset],
    ).fetchall()

    statuses = [
        r["status"] for r in conn.execute(
            "SELECT DISTINCT status FROM jobs WHERE status != 'new' ORDER BY status"
        ).fetchall()
    ]
    # Insert the employer-rejection pseudo-status right after the plain
    # 'rejected' one it's a subset of, but only once such a row actually
    # exists (same "only show what's observed" spirit as the query above).
    if "rejected" in statuses and conn.execute(
        f"SELECT 1 FROM jobs WHERE status = 'rejected' AND rejected_stage IN "
        f"({','.join('?' for _ in REJECTED_BY_EMPLOYER_STAGES)}) LIMIT 1",
        REJECTED_BY_EMPLOYER_STAGES,
    ).fetchone():
        statuses.insert(statuses.index("rejected") + 1, "rejected_by_employer")
    companies = [
        r["company"] for r in conn.execute(
            "SELECT DISTINCT company FROM jobs WHERE status != 'new' AND company IS NOT NULL ORDER BY company"
        ).fetchall()
    ]
    # Scoped to the currently selected company (if any) -- so picking a
    # company narrows the location suggestions to ones that actually occur
    # for it, instead of every location across the whole database.
    location_where = "status != 'new' AND location IS NOT NULL"
    location_params: list = []
    if company:
        location_where += " AND company = ?"
        location_params.append(company)
    raw_locations = [
        r["location"] for r in conn.execute(
            f"SELECT DISTINCT location FROM jobs WHERE {location_where}", location_params
        ).fetchall()
    ]
    locations = sorted({_normalize_city(l) for l in raw_locations if l})
    kpis = {
        "strong_matches": conn.execute(
            "SELECT COUNT(*) FROM jobs WHERE relevance_score >= ?", (STRONG_MATCH_THRESHOLD,)
        ).fetchone()[0],
        "pending_approval": conn.execute(
            "SELECT COUNT(*) FROM jobs WHERE status = 'pending_approval'"
        ).fetchone()[0],
        "letter_drafted": conn.execute(
            "SELECT COUNT(*) FROM jobs WHERE status = 'letter_drafted'"
        ).fetchone()[0],
        "new_today": conn.execute(
            "SELECT COUNT(*) FROM jobs WHERE date(scraped_at) = date('now')"
        ).fetchone()[0],
    }
    sidebar_counts = _sidebar_counts(conn)
    conn.close()

    filters = {
        "title": title,
        "company": company,
        "location": location,
        "status": status,
        "min_score": min_score,
    }
    return render_template(
        "index.html",
        jobs=jobs,
        statuses=statuses,
        companies=companies,
        locations=locations,
        score_options=SCORE_OPTIONS,
        filters=filters,
        page=page,
        total_pages=total_pages,
        total=total,
        kpis=kpis,
        sidebar_counts=sidebar_counts,
        status_meta=status_meta,
        strong_match_threshold=STRONG_MATCH_THRESHOLD,
    )


def _match_strength(value: float) -> str:
    if value >= 0.7:
        return "strong"
    if value >= 0.3:
        return "average"
    return "weak"


def _match_reasons(job: dict) -> list[dict]:
    # No made-up criteria -- only the components matching/scorer.py actually
    # uses to compute the score.
    breakdown = score_breakdown(job, load_config())
    if breakdown["excluded"]:
        return [{"label": "Title contains an exclusion word", "value": "weak"}]
    location_strength = (
        "strong" if breakdown["location_score"] >= 1.0 else ("weak" if breakdown["penalty"] < 0 else "average")
    )
    return [
        {"label": "Title/role match", "value": _match_strength(breakdown["title_score"])},
        {"label": "Location", "value": location_strength},
    ]


def _timeline(job: dict, letter: sqlite3.Row | None, interview_prep: sqlite3.Row | None) -> list[dict]:
    # Three fixed progress steps regardless of outcome, so a rejection at any
    # point in the funnel (self-rejected before sending, or the employer
    # saying no after applying or after an interview) is visually distinct --
    # each reached-state is derived from data that survives a later status
    # overwrite to 'rejected' (letter/sent_at/interview_prep rows don't get
    # deleted, unlike job['status'] itself).
    reviewed_reached = letter is not None
    sent_reached = bool(job["sent_at"])
    interview_reached = (
        interview_prep is not None or job["status"] == "interview" or job["rejected_stage"] == "interview"
    )

    steps = [
        {
            "label": "Reviewed",
            "reached": reviewed_reached,
            "submeta": (letter["generated_at"] if letter else None) or "not yet",
        },
        {
            "label": "Sent",
            "reached": sent_reached,
            "submeta": job["sent_at"] or "not yet",
        },
        {
            "label": "Interview",
            "reached": interview_reached,
            "submeta": (interview_prep["generated_at"] if interview_prep else None)
            or ("not yet" if not interview_reached else "-"),
        },
    ]
    if job["status"] == "rejected":
        steps.append({
            "label": "Rejected",
            "reached": True,
            "submeta": job["rejected_reason"] or "no reason given",
            "danger": True,
        })
    return steps


@app.route("/job/<int:job_id>")
def job_detail(job_id):
    conn = get_db()
    job = conn.execute("SELECT * FROM jobs WHERE id = ?", (job_id,)).fetchone()
    letter = conn.execute(
        "SELECT * FROM letters WHERE job_id = ? ORDER BY id DESC LIMIT 1", (job_id,)
    ).fetchone()
    interview_prep = conn.execute(
        "SELECT * FROM interview_preps WHERE job_id = ? ORDER BY id DESC LIMIT 1", (job_id,)
    ).fetchone()
    sidebar_counts = _sidebar_counts(conn)
    conn.close()
    return render_template(
        "job_detail.html",
        job=job,
        letter=letter,
        interview_prep=interview_prep,
        saved=request.args.get("saved"),
        regenerated=request.args.get("regenerated"),
        prep_generated=request.args.get("prep_generated"),
        status_meta=status_meta,
        sidebar_counts=sidebar_counts,
        match_reasons=_match_reasons(dict(job)),
        timeline=_timeline(job, letter, interview_prep),
        font_options=FONT_OPTIONS,
        selected_font=request.args.get("font", DEFAULT_FONT),
        all_projects=load_projects(),
    )


@app.route("/job/<int:job_id>/save", methods=["POST"])
def save_draft(job_id):
    final_text = request.form["final_text"]
    conn = get_db()
    letter = conn.execute(
        "SELECT id FROM letters WHERE job_id = ? ORDER BY id DESC LIMIT 1", (job_id,)
    ).fetchone()
    if letter:
        conn.execute(
            "UPDATE letters SET final_text = ?, edited = 1 WHERE id = ?",
            (final_text, letter["id"]),
        )
    else:
        # Can happen if a job was approved via the dashboard but the letter
        # generator (approval_worker.py) hasn't run yet -- without this, a
        # manually typed letter here would silently be lost.
        conn.execute(
            "INSERT INTO letters (job_id, draft, final_text, generated_at, edited) "
            "VALUES (?, ?, ?, CURRENT_TIMESTAMP, 1)",
            (job_id, final_text, final_text),
        )
    conn.execute("UPDATE jobs SET status = 'reviewed' WHERE id = ?", (job_id,))
    conn.commit()
    conn.close()
    return redirect(url_for("job_detail", job_id=job_id, saved=1))


@app.route("/job/<int:job_id>/regenerate", methods=["POST"])
def regenerate_letter(job_id):
    feedback = request.form.get("feedback", "").strip()
    project_ids = request.form.getlist("project_ids")
    if not feedback and not project_ids:
        return redirect(url_for("job_detail", job_id=job_id))

    conn = get_db()
    job = conn.execute("SELECT * FROM jobs WHERE id = ?", (job_id,)).fetchone()
    letter = conn.execute(
        "SELECT * FROM letters WHERE job_id = ? ORDER BY id DESC LIMIT 1", (job_id,)
    ).fetchone()
    conn.close()

    previous_draft = (letter["final_text"] or letter["draft"] or "") if letter else ""
    language = (letter["language"] if letter else None) or detect_language(job["description"] or job["title"] or "")

    draft = generate_letter(
        dict(job),
        cv_text=load_cv(language),
        projects=load_projects(),
        dry_run=DRY_RUN,
        language=language,
        previous_draft=previous_draft,
        feedback=feedback,
        selected_project_ids=project_ids or None,
    )

    conn = get_db()
    save_letter(conn, job_id, draft, language=language)
    conn.close()
    return redirect(url_for("job_detail", job_id=job_id, regenerated=1))


@app.route("/job/<int:job_id>/interview_prep", methods=["POST"])
def interview_prep(job_id):
    conn = get_db()
    job = conn.execute("SELECT * FROM jobs WHERE id = ?", (job_id,)).fetchone()
    letter = conn.execute(
        "SELECT * FROM letters WHERE job_id = ? ORDER BY id DESC LIMIT 1", (job_id,)
    ).fetchone()
    conn.close()

    letter_text = (letter["final_text"] or letter["draft"] or "") if letter else ""
    language = (letter["language"] if letter else None) or detect_language(job["description"] or job["title"] or "")

    content = generate_interview_prep(
        dict(job),
        cv_text=load_cv(language),
        projects=load_projects(),
        dry_run=DRY_RUN,
        language=language,
        letter_text=letter_text,
    )

    conn = get_db()
    save_interview_prep(conn, job_id, content)
    conn.close()
    return redirect(url_for("job_detail", job_id=job_id, prep_generated=1))


@app.route("/job/<int:job_id>/download.docx")
def download_docx(job_id):
    conn = get_db()
    job = conn.execute("SELECT * FROM jobs WHERE id = ?", (job_id,)).fetchone()
    letter = conn.execute(
        "SELECT * FROM letters WHERE job_id = ? ORDER BY id DESC LIMIT 1", (job_id,)
    ).fetchone()
    conn.close()
    if not letter:
        return "no draft letter for this job", 404

    from docx import Document

    from letters.document_style import build_docx_footer, build_docx_header, docx_add_bullet, render_text_blocks

    font = _font_choice()
    language = letter["language"] or detect_language(letter["final_text"] or letter["draft"] or "")
    subject = f"Re: {job['title']} ({job['company']})"
    # Textarea edits from the dashboard submit with \r\n line endings, which
    # wouldn't match the "\n\n" paragraph-break/bullet-line checks below --
    # normalize before splitting, not just for freshly AI-drafted text.
    text = (letter["final_text"] or letter["draft"] or "").replace("\r\n", "\n").replace("\r", "\n")

    doc = Document()
    doc.styles["Normal"].font.name = font["docx"]
    build_docx_header(doc.sections[0], language, subject, font["docx"])
    build_docx_footer(doc.sections[0], font["docx"])
    render_text_blocks(
        text,
        add_paragraph=lambda t: doc.add_paragraph(t),
        add_bullet=lambda t: docx_add_bullet(doc, t, font["docx"]),
    )

    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)
    return send_file(
        buf,
        as_attachment=True,
        download_name=_letter_filename(dict(job), "docx"),
        mimetype="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    )


@app.route("/job/<int:job_id>/download.pdf")
def download_pdf(job_id):
    conn = get_db()
    job = conn.execute("SELECT * FROM jobs WHERE id = ?", (job_id,)).fetchone()
    letter = conn.execute(
        "SELECT * FROM letters WHERE job_id = ? ORDER BY id DESC LIMIT 1", (job_id,)
    ).fetchone()
    conn.close()
    if not letter:
        return "no draft letter for this job", 404

    from fpdf.enums import XPos, YPos

    from letters.document_style import make_letterhead_pdf, pdf_add_bullet, render_text_blocks

    font = _font_choice()
    language = letter["language"] or detect_language(letter["final_text"] or letter["draft"] or "")
    # Textarea edits from the dashboard submit with \r\n line endings, which
    # wouldn't match the "\n\n" paragraph-break/bullet-line checks below --
    # normalize before splitting, not just for freshly AI-drafted text.
    text = (letter["final_text"] or letter["draft"] or "").replace("\r\n", "\n").replace("\r", "\n")
    # Core PDF fonts only support cp1252 -- replace characters that don't fit
    # (e.g. a single non-Latin name) instead of crashing on one stray character.
    safe_text = text.encode("cp1252", "replace").decode("cp1252")
    subject = f"Re: {job['title']} ({job['company']})".encode("cp1252", "replace").decode("cp1252")

    pdf = make_letterhead_pdf(font["pdf"], language, subject)
    pdf.add_page()

    def add_paragraph(t: str) -> None:
        pdf.set_font(font["pdf"], size=11)
        pdf.set_text_color(0, 0, 0)
        pdf.multi_cell(0, 6, t, align="L", new_x=XPos.LMARGIN, new_y=YPos.NEXT)
        pdf.ln(4)

    def add_bullet(t: str) -> None:
        pdf.set_text_color(0, 0, 0)
        pdf_add_bullet(pdf, t, font["pdf"], size=11, line_height=6)
        pdf.ln(2)

    render_text_blocks(safe_text, add_paragraph, add_bullet)

    buf = io.BytesIO(bytes(pdf.output()))
    buf.seek(0)
    return send_file(
        buf,
        as_attachment=True,
        download_name=_letter_filename(dict(job), "pdf"),
        mimetype="application/pdf",
    )


# Which rejected_stage a reject click maps to, based on the job's status right
# before the reject -- pending_approval/sent/interview are explicit funnel
# points, anything else (letter_drafted/reviewed) falls back to "review" like
# before. "applied"/"interview" are the employer saying no post-send, as
# opposed to "approval"/"review" which are Wychert deciding not to pursue it.
REJECT_STAGE_BY_STATUS = {"pending_approval": "approval", "sent": "applied", "interview": "interview"}


@app.route("/job/<int:job_id>/status", methods=["POST"])
def update_status(job_id):
    # Only updates the status in the database. Doesn't send/post anything --
    # actually sending the application is done by you, outside this dashboard.
    new_status = request.form["status"]
    if new_status not in ("sent", "rejected", "approved", "interview"):
        return "invalid status", 400
    conn = get_db()
    if new_status == "rejected":
        current = conn.execute("SELECT status FROM jobs WHERE id = ?", (job_id,)).fetchone()
        stage = REJECT_STAGE_BY_STATUS.get(current["status"] if current else None, "review")
        reason = request.form.get("reason", "").strip() or "Rejected via dashboard"
        conn.execute(
            "UPDATE jobs SET status = ?, rejected_at = CURRENT_TIMESTAMP, "
            "rejected_stage = ?, rejected_reason = ? WHERE id = ?",
            (new_status, stage, reason, job_id),
        )
    elif new_status == "sent":
        conn.execute("UPDATE jobs SET status = ?, sent_at = CURRENT_TIMESTAMP WHERE id = ?", (new_status, job_id))
    else:
        conn.execute("UPDATE jobs SET status = ? WHERE id = ?", (new_status, job_id))
    conn.commit()
    conn.close()
    # Inline status changes from the jobs list (index.html) carry the current
    # filters/page along as a raw query string so the redirect lands back on
    # the same filtered view instead of resetting to an unfiltered page 1.
    return_qs = request.form.get("return_qs", "").strip()
    if return_qs:
        return redirect(f"{url_for('index')}?{return_qs}")
    return redirect(url_for("index"))


def _cv_filename(extension: str, variant: str) -> str:
    name = SENDER["name"].replace(" ", "_")
    suffix = "_short" if variant == "short" else ""
    return f"CV_{name}{suffix}.{extension}"


@app.route("/cv")
def cv():
    conn = get_db()
    sidebar_counts = _sidebar_counts(conn)
    conn.close()
    variant = request.args.get("variant", "long")
    has_dutch_cv = CV_SHORT_PATH_NL.exists() if variant == "short" else CV_PATH_NL.exists()
    return render_template(
        "cv.html",
        font_options=FONT_OPTIONS,
        selected_font=request.args.get("font", DEFAULT_FONT),
        selected_lang=request.args.get("lang", "en"),
        selected_variant=variant,
        has_dutch_cv=has_dutch_cv,
        sidebar_counts=sidebar_counts,
    )


@app.route("/cv/download.docx")
def cv_download_docx():
    font = _font_choice()
    language = request.args.get("lang", "en")
    variant = request.args.get("variant", "long")

    doc = (
        build_short_cv_docx(language=language, font_name=font["docx"])
        if variant == "short"
        else build_cv_docx(language=language, font_name=font["docx"])
    )
    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)
    return send_file(
        buf,
        as_attachment=True,
        download_name=_cv_filename("docx", variant),
        mimetype="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    )


@app.route("/cv/download.pdf")
def cv_download_pdf():
    font = _font_choice()
    language = request.args.get("lang", "en")
    variant = request.args.get("variant", "long")

    pdf = (
        build_short_cv_pdf(language=language, font_family=font["pdf"])
        if variant == "short"
        else build_cv_pdf(language=language, font_family=font["pdf"])
    )
    buf = io.BytesIO(bytes(pdf.output()))
    buf.seek(0)
    return send_file(
        buf,
        as_attachment=True,
        download_name=_cv_filename("pdf", variant),
        mimetype="application/pdf",
    )


if __name__ == "__main__":
    app.run(debug=True, port=5000)
